"""Generates public/blank.docx, public/sample.docx and tests/fixtures/sample.docx.

Uses python-docx (dev-time only): pip install python-docx pillow
"""
import io
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw


def a4(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.orientation = WD_ORIENT.PORTRAIT
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.54)
    s.header_distance = s.footer_distance = Cm(1.25)


def korean_defaults(doc):
    # 맑은 고딕 10pt, 문단 아래 8pt, 줄 간격 1.08 (Korean Word defaults).
    styles = doc.styles.element
    rpr = styles.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    fonts = rpr.find(qn('w:rFonts'))
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme'):
        fonts.attrib.pop(qn(a), None)
    fonts.set(qn('w:ascii'), '맑은 고딕')
    fonts.set(qn('w:hAnsi'), '맑은 고딕')
    fonts.set(qn('w:eastAsia'), '맑은 고딕')
    sz = rpr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz'); rpr.append(sz)
    sz.set(qn('w:val'), '20')
    normal = doc.styles['Normal']
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.08


def blank():
    doc = Document()
    a4(doc)
    korean_defaults(doc)
    doc.core_properties.title = ''
    doc.core_properties.author = 'Oword'
    doc.save('public/blank.docx')


def picture():
    im = Image.new('RGB', (360, 180), (238, 243, 252))
    d = ImageDraw.Draw(im)
    d.rectangle([10, 10, 349, 169], outline=(43, 111, 214), width=4)
    d.ellipse([40, 40, 140, 140], fill=(242, 182, 50))
    d.polygon([(160, 150), (240, 60), (330, 150)], fill=(47, 158, 91))
    b = io.BytesIO(); im.save(b, 'PNG'); b.seek(0)
    return b


def sample():
    doc = Document()
    a4(doc)
    korean_defaults(doc)
    doc.core_properties.title = 'Oword 샘플 문서'
    doc.core_properties.author = 'Oword'
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = 'Oword 샘플 문서 — 머리글'
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run('- ')
    r = fp.add_run()
    for kind, text in (('begin', None), (None, ' PAGE '), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
        elif text == ' PAGE ':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = text; r._r.append(it)
        else:
            t = OxmlElement('w:t'); t.text = text; r._r.append(t)
    fp.add_run(' -')

    doc.add_heading('Oword 샘플 문서', level=0)
    p = doc.add_paragraph('이 문서는 ')
    p.add_run('Oword').bold = True
    p.add_run(' 편집기의 ')
    r = p.add_run('렌더링과 편집')
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(' 기능을 확인하기 위한 샘플입니다. DOCX 파일을 그대로 열고, 손대지 않은 부분은 그대로 저장합니다.')

    doc.add_heading('1. 글자 모양', level=1)
    p = doc.add_paragraph()
    p.add_run('굵게').bold = True
    p.add_run(', ')
    p.add_run('기울임').italic = True
    p.add_run(', ')
    p.add_run('밑줄').underline = True
    p.add_run(', ')
    r = p.add_run('취소선'); r.font.strike = True
    p.add_run(', H')
    r = p.add_run('2'); r.font.subscript = True
    p.add_run('O, x')
    r = p.add_run('2'); r.font.superscript = True
    p.add_run(', ')
    r = p.add_run('큰 글자 16pt'); r.font.size = Pt(16)
    p.add_run(', ')
    r = p.add_run('형광펜'); r.font.highlight_color = 7  # yellow
    p.add_run('.')

    doc.add_heading('2. 목록', level=1)
    for t in ('첫째 글머리표 항목', '둘째 글머리표 항목', '셋째 글머리표 항목'):
        doc.add_paragraph(t, style='List Bullet')
    for t in ('번호 매긴 첫 항목', '번호 매긴 둘째 항목'):
        doc.add_paragraph(t, style='List Number')

    doc.add_heading('3. 표', level=1)
    t = doc.add_table(rows=3, cols=3)
    t.style = 'Table Grid'
    data = [('항목', '수량', '비고'), ('사과', '3', '빨강'), ('배', '5', '노랑')]
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
            if i == 0:
                t.cell(i, j).paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

    doc.add_heading('4. 그림', level=1)
    doc.add_picture(picture(), width=Cm(8))
    cap = doc.add_paragraph('그림 1. 샘플 그림')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph('다음 쪽으로 넘어갑니다.')
    p.add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading('5. 정렬과 들여쓰기', level=1)
    p = doc.add_paragraph('가운데 정렬된 문단입니다.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('오른쪽 정렬된 문단입니다.'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph('첫 줄을 1cm 들여 쓴 양쪽 정렬 문단입니다. ' * 4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1)
    p = doc.add_paragraph('탭으로\t구분한\t글자')
    for i in range(6):
        doc.add_paragraph(f'채우기 문단 {i + 1}. 한 쪽을 넘는 문서에서 쪽 나누기와 머리글/바닥글이 어떻게 보이는지 확인합니다. ' * 3)
    doc.save('tests/fixtures/sample.docx')
    doc.save('public/sample.docx')


blank()
sample()
print('ok')
