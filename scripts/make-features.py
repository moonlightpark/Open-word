"""Generates tests/fixtures/features.docx and public/features.docx: a multi-page document that
exercises every feature Oword renders or must preserve. Dev-time only: pip install python-docx pillow
"""
import io
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_UNDERLINE, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NSD = (f'xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
       'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
       'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
       'xmlns:v="urn:schemas-microsoft-com:vml" '
       'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"')

LOREM = ('편집기는 문서를 열 때 모든 XML 파트를 읽어 두고, 바꾼 부분만 다시 씁니다. '
         '한글과 English가 섞인 문장, 숫자 12,345와 기호(괄호·따옴표 “ ”)도 줄바꿈과 정렬이 자연스러워야 합니다. ')


def raw(doc, xml):
    """Insert a raw WordprocessingML block before the final section properties."""
    el = parse_xml(f'<w:wrap {NSD}>{xml}</w:wrap>')[0]
    doc.element.body.sectPr.addprevious(el)
    return el


def raw_runs(p, xml):
    for el in parse_xml(f'<w:wrap {NSD}>{xml}</w:wrap>'):
        p._p.append(el)


def picture(w=360, h=180, seed=0):
    im = Image.new('RGB', (w, h), (238, 243, 252))
    d = ImageDraw.Draw(im)
    d.rectangle([6, 6, w - 7, h - 7], outline=(43, 111, 214), width=4)
    d.ellipse([30, 30, 30 + h // 2, 30 + h // 2], fill=[(242, 182, 50), (214, 64, 64)][seed % 2])
    d.polygon([(h - 20, h - 25), (w // 2, 40), (w - 30, h - 25)], fill=(47, 158, 91))
    b = io.BytesIO(); im.save(b, 'PNG'); b.seek(0)
    return b


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.54)
    s.header_distance = s.footer_distance = Cm(1.25)
    styles = doc.styles.element
    rpr = styles.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    fonts = rpr.find(qn('w:rFonts'))
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        fonts.attrib.pop(qn(a), None)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        fonts.set(qn(a), '맑은 고딕')
    sz = rpr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz'); rpr.append(sz)
    sz.set(qn('w:val'), '20')
    normal = doc.styles['Normal']
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.08


def field(p, instr, result):
    """Complex field in separate runs, as Word writes it."""
    raw_runs(p, f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
                f'<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>{result}</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>')


def headers(doc):
    s = doc.sections[0]
    s.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    s.header.paragraphs[0].text = 'Oword 기능 점검 — 홀수 쪽 머리글'
    s.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    s.even_page_header.paragraphs[0].text = 'Oword 기능 점검 — 짝수 쪽 머리글'
    for f in (s.footer, s.even_page_footer):
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('쪽 ')
        field(p, 'PAGE', '1')
        p.add_run(' / ')
        field(p, 'NUMPAGES', '1')
    s.first_page_footer.paragraphs[0].text = '표지 바닥글 (첫 쪽 전용)'
    s.first_page_footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_part(doc, name, ct, rt, xml):
    part = Part(PackURI(name), ct, xml.encode('utf-8'), doc.part.package)
    doc.part.relate_to(part, rt)


def numbering(doc):
    """A custom multi-level list (1. 가. ① i)) and a restarted instance."""
    numxml = doc.part.numbering_part.element
    lv = [('decimal', '%1.'), ('ganada', '%2.'), ('decimalEnclosedCircle', '%3'), ('lowerRoman', '%4)'), ('chosung', '%5.'),
          ('upperLetter', '%6.'), ('koreanDigital', '%7.'), ('upperRoman', '%8.'), ('decimal', '%1.%2.%9')]
    lvls = ''.join(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{f}"/><w:lvlText w:val="{t}"/><w:lvlJc w:val="left"/>'
                   f'<w:pPr><w:ind w:left="{800 + 400 * i}" w:hanging="400"/></w:pPr></w:lvl>' for i, (f, t) in enumerate(lv))
    absn = parse_xml(f'<w:abstractNum {NSD} w:abstractNumId="90"><w:multiLevelType w:val="multilevel"/>{lvls}</w:abstractNum>')
    first_num = numxml.find(qn('w:num'))
    first_num.addprevious(absn)
    numxml.append(parse_xml(f'<w:num {NSD} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    numxml.append(parse_xml(f'<w:num {NSD} w:numId="91"><w:abstractNumId w:val="90"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride></w:num>'))


def numbered(doc, text, num_id, ilvl):
    p = doc.add_paragraph(text, style='List Paragraph')
    raw_p = p._p.get_or_add_pPr()
    raw_p.append(parse_xml(f'<w:numPr {NSD}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'))
    return p


def features():
    doc = Document()
    setup(doc)
    headers(doc)
    numbering(doc)
    doc.core_properties.title = 'Oword 기능 점검 문서'
    doc.core_properties.author = 'Oword 테스트'

    # ---------------------------------------------------------------- 표지
    doc.add_paragraph('Oword 기능 점검 문서', style='Title')
    doc.add_paragraph('렌더링 · 편집 · 보존을 확인하는 여러 쪽짜리 문서', style='Subtitle')
    p = doc.add_paragraph('작성: ')
    raw_runs(p, f'<w:fldSimple {NSD} w:instr=" AUTHOR "><w:r><w:t>Oword 테스트</w:t></w:r></w:fldSimple>')
    p = doc.add_paragraph('표지 끝.')
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- 목차 (여러 문단에 걸친 필드)
    doc.add_paragraph('목차', style='TOC Heading' if 'TOC Heading' in [s.name for s in doc.styles] else 'Heading 1')
    entries = [('1. 글자 모양', '3'), ('2. 문단 모양', '4'), ('3. 목록', '5'), ('4. 표', '6')]
    for i, (t, pg) in enumerate(entries):
        begin = '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>' if i == 0 else ''
        end = '<w:r><w:fldChar w:fldCharType="end"/></w:r>' if i == len(entries) - 1 else ''
        raw(doc, f'<w:p><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9016"/></w:tabs></w:pPr>{begin}'
                 f'<w:hyperlink w:anchor="_Toc{i}" w:history="1"><w:r><w:t>{t}</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>{pg}</w:t></w:r></w:hyperlink>{end}</w:p>')
    p = doc.add_paragraph('목차 끝.')
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- 1. 글자 모양
    raw(doc, '<w:bookmarkStart w:id="0" w:name="_Toc0"/>')
    doc.add_heading('1. 글자 모양', level=1)
    raw(doc, '<w:bookmarkEnd w:id="0"/>')
    doc.add_heading('1.1 기본 모양', level=2)
    p = doc.add_paragraph()
    for text, attr in [('굵게', 'bold'), ('기울임', 'italic'), ('굵은 기울임', 'bi')]:
        r = p.add_run(text)
        r.bold = attr in ('bold', 'bi'); r.italic = attr in ('italic', 'bi')
        p.add_run(', ')
    r = p.add_run('취소선'); r.font.strike = True
    p.add_run(', ')
    r = p.add_run('이중 취소선'); r.font.double_strike = True
    p.add_run(', 위 첨자 x')
    r = p.add_run('2'); r.font.superscript = True
    p.add_run(', 아래 첨자 H')
    r = p.add_run('2'); r.font.subscript = True
    p.add_run('O, ')
    r = p.add_run('All Caps'); r.font.all_caps = True
    p.add_run(', ')
    r = p.add_run('Small Caps'); r.font.small_caps = True
    p.add_run(', ')
    r = p.add_run('자간 넓게'); r._r.get_or_add_rPr().append(parse_xml(f'<w:spacing {NSD} w:val="60"/>'))
    p.add_run(', 숨은 글자[')
    r = p.add_run('보이면 안 됨'); r.font.hidden = True
    p.add_run('].')

    doc.add_heading('1.2 밑줄 종류', level=2)
    p = doc.add_paragraph()
    for name, u in [('실선', WD_UNDERLINE.SINGLE), ('이중선', WD_UNDERLINE.DOUBLE), ('점선', WD_UNDERLINE.DOTTED),
                    ('파선', WD_UNDERLINE.DASH), ('물결선', WD_UNDERLINE.WAVY), ('굵은 선', WD_UNDERLINE.THICK), ('단어만', WD_UNDERLINE.WORDS)]:
        r = p.add_run(name); r.font.underline = u
        p.add_run('  ')

    doc.add_heading('1.3 크기와 글꼴', level=2)
    p = doc.add_paragraph()
    for s in (8, 10, 12, 16, 20, 28, 36):
        r = p.add_run(f'{s}pt '); r.font.size = Pt(s)
    p = doc.add_paragraph()
    for f in ('맑은 고딕', '바탕', '굴림', '궁서', 'Arial', 'Times New Roman', 'Courier New', 'Georgia'):
        r = p.add_run(f'{f} 가나다 Abc  '); r.font.name = f
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), f)

    doc.add_heading('1.4 색과 형광펜', level=2)
    p = doc.add_paragraph()
    for name, rgb in [('빨강', 'C00000'), ('파랑', '0070C0'), ('초록', '00B050'), ('보라', '7030A0')]:
        r = p.add_run(name + ' '); r.font.color.rgb = RGBColor.from_string(rgb)
    r = p.add_run('테마 강조 1 ')
    r._r.get_or_add_rPr().append(parse_xml(f'<w:color {NSD} w:val="4472C4" w:themeColor="accent1"/>'))
    r = p.add_run('테마 어둡게')
    r._r.get_or_add_rPr().append(parse_xml(f'<w:color {NSD} w:val="2F5496" w:themeColor="accent1" w:themeShade="BF"/>'))
    p = doc.add_paragraph()
    for name, c in [('노랑 형광펜', WD_COLOR_INDEX.YELLOW), ('초록 형광펜', WD_COLOR_INDEX.BRIGHT_GREEN), ('하늘 형광펜', WD_COLOR_INDEX.TURQUOISE)]:
        r = p.add_run(name); r.font.highlight_color = c
        p.add_run(' ')
    r = p.add_run('글자 음영')
    r._r.get_or_add_rPr().append(parse_xml(f'<w:shd {NSD} w:val="clear" w:color="auto" w:fill="D9E2F3"/>'))

    # ---------------------------------------------------------------- 2. 문단 모양
    raw(doc, '<w:bookmarkStart w:id="1" w:name="_Toc1"/>')
    h = doc.add_heading('2. 문단 모양', level=1)
    h.paragraph_format.page_break_before = True
    raw(doc, '<w:bookmarkEnd w:id="1"/>')
    doc.add_heading('2.1 정렬', level=2)
    for name, al in [('왼쪽', WD_ALIGN_PARAGRAPH.LEFT), ('가운데', WD_ALIGN_PARAGRAPH.CENTER), ('오른쪽', WD_ALIGN_PARAGRAPH.RIGHT)]:
        doc.add_paragraph(f'{name} 정렬 문단입니다.').alignment = al
    doc.add_paragraph('양쪽 정렬: ' + LOREM * 2).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph('배분 정렬 문단').alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
    doc.add_heading('2.2 들여쓰기', level=2)
    p = doc.add_paragraph('왼쪽 2cm, 오른쪽 2cm 여백: ' + LOREM); p.paragraph_format.left_indent = Cm(2); p.paragraph_format.right_indent = Cm(2)
    p = doc.add_paragraph('첫 줄 1cm 들여쓰기: ' + LOREM); p.paragraph_format.first_line_indent = Cm(1)
    p = doc.add_paragraph('첫 줄 1cm 내어쓰기: ' + LOREM); p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1)
    doc.add_heading('2.3 간격과 줄 간격', level=2)
    p = doc.add_paragraph('문단 위 12pt, 아래 12pt 간격.'); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(12)
    for name, ls in [('1줄', 1.0), ('1.5줄', 1.5), ('2줄', 2.0)]:
        p = doc.add_paragraph(f'줄 간격 {name}: ' + LOREM); p.paragraph_format.line_spacing = ls
    p = doc.add_paragraph('줄 간격 고정 20pt: ' + LOREM); p.paragraph_format.line_spacing = Pt(20)
    p = doc.add_paragraph('줄 간격 최소 30pt: ' + LOREM)
    p._p.get_or_add_pPr().append(parse_xml(f'<w:spacing {NSD} w:line="600" w:lineRule="atLeast"/>'))
    doc.add_heading('2.4 테두리, 음영, 탭', level=2)
    p = doc.add_paragraph('테두리와 음영이 있는 문단입니다. ' + LOREM)
    p._p.get_or_add_pPr().append(parse_xml(f'<w:pBdr {NSD}><w:top w:val="single" w:sz="8" w:space="4" w:color="2B6FD6"/><w:left w:val="single" w:sz="8" w:space="4" w:color="2B6FD6"/><w:bottom w:val="double" w:sz="6" w:space="4" w:color="2B6FD6"/><w:right w:val="single" w:sz="8" w:space="4" w:color="2B6FD6"/></w:pBdr>'))
    p._p.get_or_add_pPr().append(parse_xml(f'<w:shd {NSD} w:val="clear" w:color="auto" w:fill="EEF3FC"/>'))
    doc.add_paragraph('이름\t부서\t전화\n홍길동\t개발\t1234')

    # ---------------------------------------------------------------- 3. 목록
    raw(doc, '<w:bookmarkStart w:id="2" w:name="_Toc2"/>')
    doc.add_heading('3. 목록', level=1)
    raw(doc, '<w:bookmarkEnd w:id="2"/>')
    doc.add_heading('3.1 글머리표', level=2)
    for t, st in [('첫 수준 항목', 'List Bullet'), ('둘째 수준 항목', 'List Bullet 2'), ('셋째 수준 항목', 'List Bullet 3'), ('다시 첫 수준', 'List Bullet')]:
        doc.add_paragraph(t, style=st)
    doc.add_heading('3.2 여러 수준 번호', level=2)
    for t, l in [('첫째 장', 0), ('첫째 절', 1), ('첫째 항', 2), ('세부 i', 3), ('세부 ii', 3), ('둘째 항', 2), ('둘째 절', 1), ('둘째 장', 0), ('초성 수준', 4), ('한글 숫자 수준', 6)]:
        numbered(doc, t, 90, l)
    doc.add_paragraph('번호를 5부터 다시 시작하는 목록:')
    for t in ('다섯째', '여섯째'):
        numbered(doc, t, 91, 0)

    # ---------------------------------------------------------------- 4. 표
    raw(doc, '<w:bookmarkStart w:id="3" w:name="_Toc3"/>')
    h = doc.add_heading('4. 표', level=1)
    h.paragraph_format.page_break_before = True
    raw(doc, '<w:bookmarkEnd w:id="3"/>')
    doc.add_heading('4.1 머리글 행과 음영', level=2)
    t = doc.add_table(rows=4, cols=4); t.style = 'Table Grid'
    for j, v in enumerate(['구분', '1분기', '2분기', '합계']):
        c = t.cell(0, j); c.text = v; c.paragraphs[0].runs[0].bold = True
        c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {NSD} w:val="clear" w:color="auto" w:fill="D9E2F3"/>'))
    for i, row in enumerate([('매출', '120', '150', '270'), ('비용', '80', '90', '170'), ('이익', '40', '60', '100')], start=1):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
            if j: t.cell(i, j).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_heading('4.2 셀 병합과 세로 정렬', level=2)
    t = doc.add_table(rows=4, cols=4); t.style = 'Table Grid'
    a = t.cell(0, 0).merge(t.cell(0, 1)); a.text = '가로 병합 (2칸)'
    b = t.cell(1, 0).merge(t.cell(3, 0)); b.text = '세로 병합 (3줄)'
    b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c = t.cell(1, 1).merge(t.cell(2, 2)); c.text = '가로·세로 병합 (2×2)'
    t.cell(0, 2).text = 'C1'; t.cell(0, 3).text = 'D1'; t.cell(1, 3).text = 'D2'; t.cell(2, 3).text = 'D3'
    t.cell(3, 1).text = 'B4'; t.cell(3, 2).text = 'C4'
    d = t.cell(3, 3); d.text = '아래 정렬'; d.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    t.rows[3].height = Cm(1.5)
    doc.add_paragraph('텍스트 방향: 가로, 세로(위에서 아래로), 세로(아래에서 위로)')
    t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'
    for i, (text, d) in enumerate([('가로 글자', None), ('세로 쓰기 ABC', 'tbRl'), ('아래에서 위로', 'btLr')]):
        cell = t.cell(0, i); cell.text = text
        if d:
            td = OxmlElement('w:textDirection'); td.set(qn('w:val'), d)
            cell._tc.get_or_add_tcPr().append(td)
    t.rows[0].height = Cm(3)
    doc.add_heading('4.3 중첩 표와 테두리 없는 가운데 표', level=2)
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    t.cell(0, 0).text = '바깥 표의 셀'
    inner = t.cell(0, 1).add_table(rows=2, cols=2); inner.style = 'Table Grid'
    for i in range(2):
        for j in range(2):
            inner.cell(i, j).text = f'안쪽 {i + 1}-{j + 1}'
    doc.add_paragraph()
    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(2):
        for j in range(2):
            t.cell(i, j).text = f'테두리 없음 {i + 1}{j + 1}'
    doc.add_heading('4.4 여러 쪽에 걸친 긴 표', level=2)
    t = doc.add_table(rows=41, cols=3); t.style = 'Table Grid'
    for j, v in enumerate(['번호', '항목', '설명']):
        t.cell(0, j).text = v
    for i in range(1, 41):
        t.cell(i, 0).text = str(i); t.cell(i, 1).text = f'항목 {i}'; t.cell(i, 2).text = '긴 표의 행입니다. 쪽 경계를 넘습니다.'
    doc.add_paragraph()

    # ---------------------------------------------------------------- 5. 그림과 개체
    h = doc.add_heading('5. 그림과 개체', level=1)
    h.paragraph_format.page_break_before = True
    doc.add_heading('5.1 인라인 그림', level=2)
    doc.add_picture(picture(), width=Cm(8))
    doc.add_paragraph('그림 1. 인라인 그림').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('5.2 떠 있는 그림', level=2)
    rid, _ = doc.part.get_or_add_image(picture(240, 160, 1))
    p = doc.add_paragraph('이 문단에는 오른쪽에 떠 있는 그림(wp:anchor)이 있습니다. ' + LOREM)
    raw_runs(p, f'<w:r><w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="2" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
                f'<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
                f'<wp:extent cx="1440000" cy="960000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/><wp:docPr id="50" name="떠 있는 그림"/><wp:cNvGraphicFramePr/>'
                f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="50" name="float.png"/><pic:cNvPicPr/></pic:nvPicPr>'
                f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1440000" cy="960000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')
    doc.add_heading('5.3 글상자, 기호, 수식, 각주', level=2)
    p = doc.add_paragraph('글상자: ')
    txbx = ('<w:txbxContent><w:p><w:r><w:rPr><w:b/></w:rPr><w:t>글상자 제목</w:t></w:r></w:p><w:p><w:r><w:t>글상자 안의 둘째 문단</w:t></w:r></w:p></w:txbxContent>')
    raw_runs(p, f'<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="2160000" cy="720000"/><wp:docPr id="60" name="글상자 1"/>'
                f'<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2160000" cy="720000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
                f'<wps:txbx>{txbx}</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
                f'<mc:Fallback><w:pict><v:shape style="width:170pt;height:57pt"><v:textbox>{txbx}</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>')
    p = doc.add_paragraph('기호: ')
    raw_runs(p, '<w:r><w:sym w:font="Wingdings" w:char="F0FC"/></w:r><w:r><w:t xml:space="preserve"> 확인, </w:t></w:r><w:r><w:sym w:font="Symbol" w:char="F061"/></w:r><w:r><w:t xml:space="preserve"> 알파, 줄 안 하이픈: A</w:t></w:r><w:r><w:noBreakHyphen/></w:r><w:r><w:t>B</w:t></w:r>')
    p = doc.add_paragraph('수식: ')
    raw_runs(p, '<m:oMath><m:r><m:t>x</m:t></m:r><m:sSup><m:e><m:r><m:t>y</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup><m:r><m:t>=1</m:t></m:r></m:oMath>')
    p = doc.add_paragraph('각주가 달린 문장')
    raw_runs(p, '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="1"/></w:r><w:r><w:t>입니다.</w:t></w:r>')
    add_part(doc, '/word/footnotes.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml', RT.FOOTNOTES,
             f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes {NSD}>'
             '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
             '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
             '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> 각주 내용입니다.</w:t></w:r></w:p></w:footnote></w:footnotes>')

    # ---------------------------------------------------------------- 6. 필드와 링크
    doc.add_heading('6. 필드와 링크', level=1)
    p = doc.add_paragraph('외부 링크: ')
    rid = doc.part.relate_to('https://github.com/moonlightpark/Open-word', RT.HYPERLINK, is_external=True)
    raw_runs(p, f'<w:hyperlink r:id="{rid}" w:history="1"><w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr><w:t>Oword 저장소</w:t></w:r></w:hyperlink>')
    p.add_run(', 문서 안 링크: ')
    raw_runs(p, '<w:hyperlink w:anchor="_Toc3" w:history="1"><w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr><w:t>4. 표로 가기</w:t></w:r></w:hyperlink>')
    p = doc.add_paragraph('이 쪽 번호는 ')
    field(p, 'PAGE', '7')
    p.add_run(', 전체 쪽 수는 ')
    field(p, 'NUMPAGES', '9')
    p.add_run('입니다.')

    # ---------------------------------------------------------------- 7. 변경 내용, 메모, 콘텐츠 컨트롤
    doc.add_heading('7. 변경 내용과 메모', level=1)
    p = doc.add_paragraph('변경 내용 추적: ')
    raw_runs(p, '<w:ins w:id="101" w:author="검토자" w:date="2026-10-16T00:00:00Z"><w:r><w:t>삽입된 글자</w:t></w:r></w:ins><w:r><w:t xml:space="preserve"> / </w:t></w:r>'
                '<w:del w:id="102" w:author="검토자" w:date="2026-10-16T00:00:00Z"><w:r><w:delText>삭제된 글자</w:delText></w:r></w:del><w:r><w:t xml:space="preserve"> 끝.</w:t></w:r>')
    p = doc.add_paragraph()
    raw_runs(p, '<w:commentRangeStart w:id="0"/><w:r><w:t>메모가 달린 글자</w:t></w:r><w:commentRangeEnd w:id="0"/><w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr><w:commentReference w:id="0"/></w:r><w:r><w:t xml:space="preserve"> 뒤의 글자.</w:t></w:r>')
    add_part(doc, '/word/comments.xml', 'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml', RT.COMMENTS,
             f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:comments {NSD}><w:comment w:id="0" w:author="검토자" w:date="2026-10-16T00:00:00Z" w:initials="검"><w:p><w:r><w:t>메모 내용</w:t></w:r></w:p></w:comment></w:comments>')
    p = doc.add_paragraph('콘텐츠 컨트롤: ')
    raw_runs(p, '<w:sdt><w:sdtPr><w:alias w:val="이름"/><w:tag w:val="name"/></w:sdtPr><w:sdtContent><w:r><w:t>홍길동</w:t></w:r></w:sdtContent></w:sdt><w:r><w:t xml:space="preserve"> 님.</w:t></w:r>')

    # ---------------------------------------------------------------- 8. 쪽과 구역
    doc.add_heading('8. 쪽과 구역', level=1)
    doc.add_paragraph('다음 문단 뒤에 구역 나누기(다음 쪽)가 있습니다.')
    p = doc.add_paragraph('구역 1의 마지막 문단.')
    sect = parse_xml(f'<w:sectPr {NSD}><w:type w:val="nextPage"/><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>')
    p._p.get_or_add_pPr().append(sect)
    doc.add_heading('9. 긴 본문', level=1)
    for i in range(24):
        doc.add_paragraph(f'{i + 1}. ' + LOREM * 3)
    doc.add_paragraph('문서 끝.')
    doc.save('tests/fixtures/features.docx')
    doc.save('public/features.docx')


features()
print('ok')
